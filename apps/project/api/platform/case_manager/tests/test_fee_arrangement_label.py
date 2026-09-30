"""
«Modalidad del contrato» pasa a «Esquema de honorarios», y la opcion
«Modalidad de pago» se muestra como «Honorarios fijos (abonos)».

Solo cambia lo que se **ve**. El valor guardado sigue siendo
`'Modalidad de pago'`: hay datos y una importacion que dependen de el.
"""

from io import BytesIO
from io import StringIO

from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import translation
from docx import Document

from ..choices import Mandate
from ..models import CaseFinanceModel
from ..reports import client_report, crm_report
from .test_access import login_as, make_user
from .test_notes import make_case
from .test_public_access import identificarse

VIEJAS = ('Modalidad del contrato', 'Contract modality', 'contract modality')


def docx_text(response):
    document = Document(BytesIO(response.content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return '\n'.join(parts)


class FeeArrangementLabelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('setup_case_manager_group', stdout=StringIO())
        cls.case = make_case()
        CaseFinanceModel.objects.create(
            case=cls.case, mandate=Mandate.PAYMENT, agreed_fee=2_000_000,
            paid_amount=500_000,
        )

    def setUp(self):
        make_user('abogada', gestor=True)
        login_as(self.client, 'abogada')

    def test_el_valor_guardado_no_cambia(self):
        self.assertEqual(Mandate.PAYMENT.value, 'Modalidad de pago')
        self.assertEqual(
            CaseFinanceModel.objects.get(case=self.case).mandate,
            'Modalidad de pago',
        )

    def test_los_nombres_en_ingles(self):
        with translation.override('en'):
            self.assertEqual(str(Mandate.PAYMENT.label), 'Fixed fees (installments)')
            self.assertEqual(
                str(CaseFinanceModel._meta.get_field('mandate').verbose_name),
                'fee arrangement',
            )

    def test_los_nombres_en_espanol(self):
        with translation.override('es'):
            self.assertEqual(str(Mandate.PAYMENT.label), 'Honorarios fijos (abonos)')
            self.assertEqual(
                str(CaseFinanceModel._meta.get_field('mandate').verbose_name),
                'esquema de honorarios',
            )

    def test_el_formulario_del_asunto(self):
        url = reverse('case_manager:gestor_case_update', args=[self.case.pk])

        en = self.client.get(url, HTTP_ACCEPT_LANGUAGE='en').content.decode()
        es = self.client.get(url, HTTP_ACCEPT_LANGUAGE='es').content.decode()

        self.assertIn('Fee arrangement', en)
        self.assertIn('Fixed fees (installments)', en)
        self.assertIn('Esquema de honorarios', es)
        self.assertIn('Honorarios fijos (abonos)', es)
        # El `value` de la opcion sigue siendo el de siempre.
        self.assertIn('value="Modalidad de pago"', en)
        for viejo in VIEJAS:
            self.assertNotIn(viejo, en)
            self.assertNotIn(viejo, es)

    def test_el_portal_del_cliente(self):
        self.client.logout()

        en = identificarse(self.client)
        self.assertContains(en, 'Fee arrangement')
        self.assertContains(en, 'Fixed fees (installments)')
        for viejo in VIEJAS:
            self.assertNotContains(en, viejo)

    def test_el_portal_en_espanol(self):
        self.client.logout()
        identificarse(self.client)
        es = self.client.get(
            reverse('case_manager:public_query'), HTTP_ACCEPT_LANGUAGE='es'
        )
        self.assertContains(es, 'Esquema de honorarios')
        self.assertContains(es, 'Honorarios fijos (abonos)')

    def test_la_ficha_del_cliente(self):
        url = reverse('case_manager:gestor_client_detail',
                      args=[self.case.client.pk])
        html = self.client.get(url).content.decode()
        self.assertIn('Fee arrangement', html)
        self.assertIn('Fixed fees (installments)', html)

    def test_los_informes_word(self):
        with translation.override('es'):
            ficha = docx_text(client_report(self.case.client))
            crm = docx_text(crm_report())
        for texto in (ficha, crm):
            self.assertIn('Esquema de honorarios', texto)
            self.assertNotIn('Modalidad', texto.replace('Modalidad de pago', ''))
        self.assertIn('Honorarios fijos (abonos)', ficha)
