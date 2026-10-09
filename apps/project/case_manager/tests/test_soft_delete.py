"""Borrado lógico: permisos, auditoría, consultas e identidad conservada."""
from io import StringIO
from unittest.mock import patch

from auditlog.models import LogEntry
from django.contrib import admin
from django.contrib.messages.storage.fallback import FallbackStorage
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import Client, RequestFactory, TestCase
from django.urls import reverse

from ..admin import DeletedFilter
from ..choices import Mandate, Service
from ..forms import ClientForm, PublicCaseQueryForm
from ..models import CaseFinanceModel, CaseModel, CaseNoteModel, ClientModel, PazYSalvoDocumentModel
from .. import portal_otp
from .test_access import login_as, make_user


class SoftDeleteTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.actor = make_user('delete_manager', gestor=True, staff=True)
        cls.person = ClientModel.objects.create(identification='11223344', full_name='Cliente secreto')
        cls.case = CaseModel.objects.create(client=cls.person, service=Service.JUDICIAL, case_number='SECRET-123')
        cls.second = CaseModel.objects.create(client=cls.person, service=Service.ADMINISTRATIVE)
        cls.finance = CaseFinanceModel.objects.create(case=cls.case, mandate=Mandate.PAYMENT,
            agreed_fee=100, payment_history=[{'kind': 'expected', 'amount': 30, 'date': '2026-12-01'}])
        cls.note = CaseNoteModel.objects.create(case=cls.case, title='Secreta', body='Privada')

    def setUp(self):
        login_as(self.client, self.actor.username)

    def url(self, name, obj=None):
        return reverse('case_manager:' + name, args=[obj.pk] if obj else [])

    def delete_client(self):
        return self.client.post(self.url('gestor_client_delete', self.person))

    def test_cascade_managers_and_audit_actor(self):
        self.assertEqual(self.delete_client().status_code, 302)
        self.person.refresh_from_db()
        self.assertFalse(self.person.is_active)
        self.assertEqual(self.person.deleted_by, self.actor)
        self.assertIsNotNone(self.person.deleted_at)
        self.assertFalse(ClientModel.objects.exists())
        self.assertEqual(ClientModel.all_objects.count(), 1)
        self.assertEqual(CaseModel.all_objects.count(), 2)
        self.assertFalse(self.person.cases.all().exists())
        self.assertEqual(ClientModel._meta.default_manager_name, 'objects')
        self.assertEqual(CaseModel._meta.base_manager_name, 'all_objects')
        for obj in [self.person, self.case, self.second]:
            obj.refresh_from_db()
            self.assertEqual(obj.deleted_at, self.person.deleted_at)
            self.assertEqual(obj.deleted_by, self.actor)
            self.assertFalse(obj.is_active)
            entry = LogEntry.objects.get_for_object(obj).filter(action=LogEntry.Action.UPDATE).latest('timestamp')
            self.assertEqual(entry.actor, self.actor)
            self.assertIn('deleted_at', entry.changes_dict)

    def test_single_case_and_reverse_manager(self):
        response = self.client.post(self.url('gestor_case_delete', self.case))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(list(self.person.cases.all()), [self.second])
        self.assertTrue(ClientModel.objects.filter(pk=self.person.pk).exists())
        response = self.client.get(self.url('gestor_client_list'))
        self.assertEqual(response.context['clients'][0].case_count, 1)
        detail = self.client.get(self.url('gestor_client_detail', self.person))
        self.assertEqual(list(detail.context['cases']), [self.second])
        self.assertFalse(CaseFinanceModel.objects.exists())
        self.assertFalse(CaseNoteModel.objects.exists())

    def test_permissions_post_only_and_csrf(self):
        for name, obj in [('gestor_case_delete', self.case), ('gestor_client_delete', self.person)]:
            url = self.url(name, obj)
            self.assertEqual(self.client.get(url).status_code, 405)
            secured = Client(enforce_csrf_checks=True)
            login_as(secured, self.actor.username)
            self.assertEqual(secured.post(url).status_code, 403)
            page = secured.get(self.url('gestor_client_list'))
            token = secured.cookies['csrftoken'].value
            self.assertEqual(secured.post(url, HTTP_X_CSRFTOKEN=token).status_code, 302)
        stranger = make_user('stranger')
        login_as(self.client, stranger.username)
        self.assertEqual(self.delete_client().status_code, 404)
        self.assertEqual(self.client.post(self.url('gestor_case_delete', self.case)).status_code, 404)

    def test_deleted_disappears_from_views_reports_and_finances(self):
        self.delete_client()
        for name in ['gestor_client_list', 'gestor_case_list', 'gestor_dashboard']:
            response = self.client.get(self.url(name))
            self.assertEqual(response.status_code, 200)
            self.assertNotContains(response, 'Cliente Secreto')
            self.assertNotContains(response, 'SECRET-123')
        dashboard = self.client.get(self.url('gestor_dashboard'))
        self.assertEqual(dashboard.context['upcoming_total'], 0)
        self.assertFalse(dashboard.context['debtors'])
        self.assertFalse(dashboard.context['expectations'])
        chart = self.client.get(self.url('gestor_dashboard'), HTTP_X_REQUESTED_WITH='fetch')
        self.assertEqual(chart.status_code, 200)
        self.assertNotContains(chart, 'SECRET-123')
        for name, obj in [('gestor_client_detail', self.person), ('gestor_client_report', self.person),
                          ('gestor_case_update', self.case), ('gestor_case_report', self.case),
                          ('paz_y_salvo', self.case), ('paz_y_salvo_download', self.case)]:
            self.assertEqual(self.client.get(self.url(name, obj)).status_code, 404, name)
        note_url = reverse('case_manager:gestor_case_note_visibility', args=[self.case.pk, self.note.pk])
        self.assertEqual(self.client.post(note_url, {'visible': '1'}).status_code, 404)
        response = self.client.get(self.url('gestor_crm_report'))
        self.assertEqual(response.status_code, 200)
        from docx import Document
        from io import BytesIO
        content = b''.join(response.streaming_content) if response.streaming else response.content
        document = Document(BytesIO(content))
        text = ' '.join(p.text for p in document.paragraphs)
        text += ' '.join(c.text for t in document.tables for row in t.rows for c in row.cells)
        self.assertNotIn('SECRET-123', text)
        self.assertNotIn('Cliente Secreto', text)

    def test_portal_and_otp_treat_deleted_like_missing(self):
        request = RequestFactory().get('/')
        request.session = self.client.session
        with patch('apps.project.case_manager.emails.send_access_code', return_value=1):
            portal_otp.issue(request, self.person)
        self.delete_client()
        for number in [self.person.identification, '99999999999']:
            form = PublicCaseQueryForm({'identification': number})
            self.assertTrue(form.is_valid())
            self.assertIsNone(form.get_client())
        self.assertFalse(portal_otp.issue(request, self.person))
        self.assertFalse(portal_otp.verify(request, '123456'))
        self.client.logout()
        responses = [self.client.post(self.url('public_query'), {'identification': n})
                     for n in [self.person.identification, '99999999999']]
        self.assertEqual(responses[0].status_code, responses[1].status_code)
        from ..forms import UNKNOWN_IDENTIFICATION
        for response in responses:
            self.assertContains(response, str(UNKNOWN_IDENTIFICATION), status_code=400)

    def test_document_is_reserved_and_restore_offer_works(self):
        self.delete_client()
        data = {'identification': self.person.identification, 'full_name': 'Otro nombre'}
        form = ClientForm(data)
        self.assertFalse(form.is_valid())
        self.assertEqual(form.deleted_client, self.person)
        self.assertIn('eliminado', str(form.errors))
        page = self.client.post(self.url('gestor_client_create'), data)
        self.assertContains(page, self.url('gestor_client_restore', self.person))
        with self.assertRaises(IntegrityError), transaction.atomic():
            ClientModel.objects.create(**data)
        self.assertEqual(self.client.post(self.url('gestor_client_restore', self.person)).status_code, 302)
        self.person.refresh_from_db()
        self.assertIsNone(self.person.deleted_at)
        self.assertIsNone(self.person.deleted_by)
        self.assertFalse(self.person.is_active)
        self.assertEqual(ClientModel.all_objects.count(), 1)

    def test_admin_lists_filters_and_restores_without_reactivation(self):
        self.delete_client()
        request = RequestFactory().post('/')
        request.user = self.actor
        request.session = self.client.session
        request._messages = FallbackStorage(request)
        for model, obj in [(ClientModel, self.person), (CaseModel, self.case)]:
            model_admin = admin.site._registry[model]
            self.assertIn(DeletedFilter, model_admin.list_filter)
            self.assertTrue(model_admin.get_queryset(request).filter(pk=obj.pk).exists())
            model_admin.restore_selected(request, model.all_objects.filter(pk=obj.pk))
            obj.refresh_from_db()
            self.assertIsNone(obj.deleted_at)
            self.assertIsNone(obj.deleted_by)
            self.assertFalse(obj.is_active)

    def test_restoring_case_under_deleted_client_keeps_it_hidden(self):
        self.delete_client()
        self.case.refresh_from_db()
        self.case.restore()
        self.assertTrue(CaseModel.all_objects.filter(pk=self.case.pk).exists())
        self.assertFalse(CaseModel.objects.filter(pk=self.case.pk).exists())
        self.assertFalse(CaseFinanceModel.objects.exists())

    def test_certified_letter_and_note_remain_inaccessible_after_delete(self):
        from django.utils import timezone
        from ..emails import send_case_note
        from .. import paz_y_salvo
        doc = PazYSalvoDocumentModel.objects.create(case=self.case,
            authorized_at=timezone.now(), status='CERTIFIED', idempotency_key='deleted-letter',
            public_copy_file='private.pdf')
        self.delete_client()
        self.assertEqual(self.client.get(self.url('paz_y_salvo', self.case)).status_code, 404)
        self.assertEqual(self.client.get(self.url('paz_y_salvo_download', self.case)).status_code, 404)
        self.assertFalse(send_case_note(self.note))
        doc.status = 'PENDING'
        doc.save()
        with patch('apps.project.case_manager.gea_client.issue') as issue:
            self.assertEqual(paz_y_salvo.certify(doc.pk).status, 'PENDING')
            issue.assert_not_called()

    def test_import_and_seed_skip_reserved_documents(self):
        import json
        import tempfile
        from pathlib import Path
        self.delete_client()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'import.json'
            path.write_text(json.dumps({self.person.identification: {'n': 'Reimportado'}}))
            call_command('import_propdemo', str(path), stdout=StringIO(), stderr=StringIO())
        module = 'apps.project.case_manager.management.commands.seed_gestor_demo.'
        with patch(module + 'handmade_clients', return_value=[{'identification': self.person.identification}]), \
             patch(module + 'generated_clients', return_value=[]), \
             patch(module + 'document_clients', return_value=[]):
            call_command('seed_gestor_demo', stdout=StringIO())
        self.person.refresh_from_db()
        self.assertIsNotNone(self.person.deleted_at)
        self.assertEqual(ClientModel.all_objects.count(), 1)
        self.assertEqual(CaseModel.all_objects.count(), 2)
